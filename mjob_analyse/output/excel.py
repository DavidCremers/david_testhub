"""Excel export module voor gedetailleerde rapportage."""

from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, LineChart, Reference

from ..config.settings import Config
from ..utils.logging import get_logger

logger = get_logger("output.excel")


class ExcelExport:
    """
    Exporteert analyseresultaten naar Excel met meerdere tabbladen.

    Bevat formatting, kleuren voor afwijkingen, en optionele grafieken.
    """

    # Kleuren
    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    WARNING_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    SUCCESS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    NEUTRAL_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    def __init__(self, config: Config):
        """
        Initialiseer de Excel exporter.

        Args:
            config: Configuratie met output instellingen
        """
        self.config = config
        self.workbook: Optional[Workbook] = None
        self._current_row = 1

    def export(
        self,
        analysis_results: dict,
        output_path: Path,
        name_a: str = "Systeem A",
        name_b: str = "Systeem B",
    ) -> Path:
        """
        Exporteer alle analyseresultaten naar Excel.

        Args:
            analysis_results: Dictionary met alle analyseresultaten
            output_path: Pad voor output bestand
            name_a: Naam van eerste systeem
            name_b: Naam van tweede systeem

        Returns:
            Pad naar gegenereerd bestand
        """
        logger.info(f"Start Excel export naar: {output_path}")

        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)  # Verwijder default sheet

        # Samenvatting tabblad
        self._create_summary_sheet(analysis_results, name_a, name_b)

        # Matches tabblad
        if "matches" in analysis_results:
            self._create_matches_sheet(
                analysis_results["matches"], name_a, name_b
            )

        # Prijsafwijkingen tabblad
        if "prijs" in analysis_results:
            self._create_price_sheet(analysis_results["prijs"], name_a, name_b)

        # Ontbrekende ingrepen tabbladen
        if "compleetheid" in analysis_results:
            self._create_completeness_sheets(
                analysis_results["compleetheid"], name_a, name_b
            )

        # Bouwdeel samenvatting
        self._create_building_part_sheet(analysis_results, name_a, name_b)

        # Cashflow vergelijking
        if "financieel" in analysis_results:
            self._create_cashflow_sheet(
                analysis_results["financieel"], name_a, name_b
            )

        # Opslaan
        output_path.parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(output_path)

        logger.info(f"Excel export voltooid: {output_path}")
        return output_path

    def _create_summary_sheet(
        self, results: dict, name_a: str, name_b: str
    ) -> None:
        """Maak samenvatting tabblad."""
        ws = self.workbook.create_sheet("Samenvatting")

        # Titel
        ws["A1"] = "MJOB Analyse Rapport"
        ws["A1"].font = Font(size=16, bold=True)
        ws.merge_cells("A1:D1")

        ws["A3"] = f"Vergelijking: {name_a} vs {name_b}"
        ws["A3"].font = Font(size=12)

        row = 5

        # Match statistieken
        if "match_quality" in results:
            ws.cell(row=row, column=1, value="Match Kwaliteit").font = Font(bold=True)
            row += 1

            mq = results["match_quality"]
            for key, value in mq.items():
                ws.cell(row=row, column=1, value=key.replace("_", " ").title())
                ws.cell(row=row, column=2, value=value)
                row += 1
            row += 1

        # Financieel overzicht
        if "financieel" in results:
            ws.cell(row=row, column=1, value="Financieel Overzicht").font = Font(bold=True)
            row += 1

            fin = results["financieel"]
            for key in [f"totaal_{name_a}", f"totaal_{name_b}"]:
                if key in fin:
                    totaal = fin[key]
                    ws.cell(row=row, column=1, value=totaal.get("naam", key))
                    ws.cell(
                        row=row, column=2, value=totaal.get("totaal_over_horizon", 0)
                    )
                    ws.cell(row=row, column=2).number_format = "€#,##0"
                    row += 1

            if "verschil_totaal" in fin:
                ws.cell(row=row, column=1, value="Verschil")
                ws.cell(row=row, column=2, value=fin["verschil_totaal"]["absoluut"])
                ws.cell(row=row, column=2).number_format = "€#,##0"
                row += 1

        # Auto-fit kolommen
        self._autofit_columns(ws)

    def _create_matches_sheet(
        self, matches_df: pd.DataFrame, name_a: str, name_b: str
    ) -> None:
        """Maak tabblad met alle matches."""
        if matches_df.empty:
            return

        ws = self.workbook.create_sheet("Alle Matches")

        # Schrijf data
        self._write_dataframe(ws, matches_df)

        # Conditional formatting voor confidence
        self._apply_confidence_formatting(ws, matches_df)

        self._autofit_columns(ws)

    def _create_price_sheet(
        self, price_results: dict, name_a: str, name_b: str
    ) -> None:
        """Maak tabblad met prijsafwijkingen."""
        ws = self.workbook.create_sheet("Prijsafwijkingen")

        # Statistieken bovenaan
        row = 1
        ws.cell(row=row, column=1, value="Prijsanalyse Statistieken").font = Font(
            bold=True
        )
        row += 1

        stats = price_results.get("statistieken", {})
        for key, value in stats.items():
            ws.cell(row=row, column=1, value=key.replace("_", " ").title())
            ws.cell(row=row, column=2, value=value)
            row += 1

        row += 2

        # Detail data
        detail_df = price_results.get("detail_data")
        if detail_df is not None and not detail_df.empty:
            # Filter alleen afwijkingen
            if "Is_Prijs_Uitschieter" in detail_df.columns:
                outliers = detail_df[detail_df["Is_Prijs_Uitschieter"]]
            else:
                outliers = detail_df

            ws.cell(row=row, column=1, value="Prijsuitschieters").font = Font(bold=True)
            row += 1

            self._write_dataframe(ws, outliers, start_row=row)

            # Kleur afwijkingen
            self._color_outlier_rows(ws, outliers, start_row=row)

        self._autofit_columns(ws)

    def _create_completeness_sheets(
        self, completeness_results: dict, name_a: str, name_b: str
    ) -> None:
        """Maak tabbladen voor ontbrekende ingrepen."""
        # Ontbrekend in A
        key_a = f"ontbrekend_in_{name_a}"
        if key_a in completeness_results:
            df_a = completeness_results[key_a]
            if isinstance(df_a, pd.DataFrame) and not df_a.empty:
                ws = self.workbook.create_sheet(f"Ontbreekt in {name_a}"[:31])
                self._write_dataframe(ws, df_a)
                self._autofit_columns(ws)

        # Ontbrekend in B
        key_b = f"ontbrekend_in_{name_b}"
        if key_b in completeness_results:
            df_b = completeness_results[key_b]
            if isinstance(df_b, pd.DataFrame) and not df_b.empty:
                ws = self.workbook.create_sheet(f"Ontbreekt in {name_b}"[:31])
                self._write_dataframe(ws, df_b)
                self._autofit_columns(ws)

    def _create_building_part_sheet(
        self, results: dict, name_a: str, name_b: str
    ) -> None:
        """Maak tabblad met samenvatting per bouwdeel."""
        ws = self.workbook.create_sheet("Per Bouwdeel")

        row = 1

        # Financiële breakdown
        if "financieel" in results:
            fin = results["financieel"]

            for system_name in [name_a, name_b]:
                breakdown_key = f"breakdown_{system_name}"
                if breakdown_key in fin:
                    ws.cell(
                        row=row, column=1, value=f"Kosten per bouwdeel - {system_name}"
                    ).font = Font(bold=True)
                    row += 1

                    # Headers
                    headers = ["Bouwdeel", "Aantal Ingrepen", "Totaal", "% van Totaal"]
                    for col, header in enumerate(headers, 1):
                        cell = ws.cell(row=row, column=col, value=header)
                        cell.fill = self.HEADER_FILL
                        cell.font = self.HEADER_FONT
                    row += 1

                    # Data
                    breakdown = fin[breakdown_key]
                    for bouwdeel, data in sorted(
                        breakdown.items(),
                        key=lambda x: x[1]["totaal_over_horizon"],
                        reverse=True,
                    ):
                        ws.cell(row=row, column=1, value=bouwdeel)
                        ws.cell(row=row, column=2, value=data["aantal_ingrepen"])
                        ws.cell(row=row, column=3, value=data["totaal_over_horizon"])
                        ws.cell(row=row, column=3).number_format = "€#,##0"
                        ws.cell(
                            row=row, column=4, value=data["percentage_van_totaal"] / 100
                        )
                        ws.cell(row=row, column=4).number_format = "0.0%"
                        row += 1

                    row += 2

        self._autofit_columns(ws)

    def _create_cashflow_sheet(
        self, financial_results: dict, name_a: str, name_b: str
    ) -> None:
        """Maak tabblad met cashflow vergelijking."""
        ws = self.workbook.create_sheet("Cashflow")

        cf_key = "cashflow_vergelijking"
        if cf_key not in financial_results:
            return

        cf_df = financial_results[cf_key]
        if cf_df.empty:
            return

        # Schrijf data
        self._write_dataframe(ws, cf_df)

        # Voeg grafiek toe
        self._add_cashflow_chart(ws, cf_df, name_a, name_b)

        self._autofit_columns(ws)

    def _write_dataframe(
        self, ws, df: pd.DataFrame, start_row: int = 1
    ) -> None:
        """Schrijf DataFrame naar worksheet."""
        # Headers
        for col, header in enumerate(df.columns, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT

        # Data
        for row_idx, row in enumerate(df.itertuples(index=False), start_row + 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx)

                # Handle pandas types
                if pd.isna(value):
                    cell.value = None
                elif isinstance(value, (pd.Timestamp,)):
                    cell.value = value.to_pydatetime()
                else:
                    cell.value = value

                # Number formatting
                if isinstance(value, float):
                    if "prijs" in str(df.columns[col_idx - 1]).lower():
                        cell.number_format = "€#,##0.00"
                    elif "percentage" in str(df.columns[col_idx - 1]).lower():
                        cell.number_format = "0.0%"

    def _apply_confidence_formatting(self, ws, df: pd.DataFrame) -> None:
        """Pas kleuren toe op basis van confidence score."""
        if "Match_Confidence" not in df.columns:
            return

        conf_col = list(df.columns).index("Match_Confidence") + 1

        for row in range(2, len(df) + 2):
            cell = ws.cell(row=row, column=conf_col)
            try:
                value = float(cell.value) if cell.value else 0
                if value >= 0.8:
                    cell.fill = self.SUCCESS_FILL
                elif value >= 0.6:
                    cell.fill = self.NEUTRAL_FILL
                else:
                    cell.fill = self.WARNING_FILL
            except (ValueError, TypeError):
                pass

    def _color_outlier_rows(
        self, ws, df: pd.DataFrame, start_row: int
    ) -> None:
        """Kleur rijen met grote afwijkingen."""
        for row_idx in range(start_row + 1, start_row + 1 + len(df)):
            for col in range(1, len(df.columns) + 1):
                ws.cell(row=row_idx, column=col).fill = self.WARNING_FILL

    def _add_cashflow_chart(
        self, ws, df: pd.DataFrame, name_a: str, name_b: str
    ) -> None:
        """Voeg cashflow grafiek toe."""
        chart = LineChart()
        chart.title = "Jaarlijkse Cashflow Vergelijking"
        chart.style = 10
        chart.y_axis.title = "Kosten (€)"
        chart.x_axis.title = "Jaar"

        # Data references
        data_end_row = len(df) + 1

        # Kosten A
        kosten_a_col = None
        kosten_b_col = None
        for idx, col in enumerate(df.columns, 1):
            if f"Kosten_{name_a}" in col:
                kosten_a_col = idx
            if f"Kosten_{name_b}" in col:
                kosten_b_col = idx

        if kosten_a_col and kosten_b_col:
            data = Reference(ws, min_col=kosten_a_col, min_row=1, max_row=data_end_row)
            chart.add_data(data, titles_from_data=True)

            data = Reference(ws, min_col=kosten_b_col, min_row=1, max_row=data_end_row)
            chart.add_data(data, titles_from_data=True)

            categories = Reference(ws, min_col=1, min_row=2, max_row=data_end_row)
            chart.set_categories(categories)

            ws.add_chart(chart, "H2")

    def _autofit_columns(self, ws) -> None:
        """Pas kolombreedte automatisch aan."""
        from openpyxl.utils import get_column_letter

        for col_idx, column_cells in enumerate(ws.columns, 1):
            try:
                length = max(
                    len(str(cell.value or "")) for cell in column_cells
                    if hasattr(cell, 'value')
                )
                adjusted_width = min(length + 2, 50)
                ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width
            except (ValueError, AttributeError):
                # Skip problematic columns
                pass
